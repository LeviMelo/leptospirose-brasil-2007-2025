# -*- coding: utf-8 -*-
"""Monta a carta de resposta a secretaria da RESS.

Uma carta de resposta so e util se cada item da secretaria puder ser conferido
sem procurar: por isso a estrutura e uma linha por item, com o que foi pedido e
o que foi feito, nomeando onde. O texto vem daqui e nao de um documento
editado a mao pela mesma razao que o manuscrito: para nao divergir do que o
gerador produz.

Saida: carta_resposta_RESS-2026-1462.docx
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

AQUI = Path(__file__).resolve().parent
SAIDA = AQUI / "carta_resposta_RESS-2026-1462.docx"

FONTE = "Times New Roman"
CORPO = Pt(12)
REPO = "https://github.com/LeviMelo/leptospirose-brasil-2007-2025"


def p(doc, texto, negrito_inicial=False, espaco_antes=0, espaco_depois=8,
      recuo=0.0):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = par.paragraph_format
    pf.space_before = Pt(espaco_antes)
    pf.space_after = Pt(espaco_depois)
    pf.line_spacing = 1.0
    if recuo:
        pf.left_indent = Cm(recuo)
    for pedaco in re.split(r"(\*\*[^*]+\*\*)", texto):
        if not pedaco:
            continue
        if pedaco.startswith("**"):
            r = par.add_run(pedaco[2:-2]); r.bold = True
        else:
            r = par.add_run(pedaco)
        r.font.name = FONTE
        r.font.size = CORPO
    return par


# Um item por solicitacao da secretaria, na ordem em que ela os numerou.
ITENS = [
    ("Não existe símbolo universal “H” para a proporção de casos internados.",
     "O símbolo foi eliminado do título, do resumo, do texto, das tabelas e "
     "das três figuras, que foram regeradas a partir dos seus programas para "
     "que nenhum eixo ou legenda o conservasse. A variável é nomeada por "
     "extenso — “proporção de casos confirmados internados” — em todas as "
     "ocorrências."),

    ("Seguir o modelo da RESS, observando a posição do quadro de aspectos "
     "éticos.",
     "O quadro passou para logo depois das palavras-chave e recebeu a norma "
     "que dispensa a apreciação ética (parágrafo único do artigo 1º da "
     "Resolução do Conselho Nacional de Saúde nº 510, de 7 de abril de 2016). "
     "O documento foi reformatado conforme o modelo: Times New Roman 12, "
     "espaçamento simples e texto alinhado à esquerda, com a ordem de seções e "
     "de material final ali prevista. O resumo tem 242 palavras e o texto da "
     "introdução ao final da discussão, 3.274, dentro dos limites de 250 e "
     "3.500."),

    ("Trata-se de análise de séries temporais; retirar “estudo ecológico” do "
     "título e das demais seções.",
     "O título passou a “Amplitude de detecção e letalidade notificada da "
     "leptospirose: análise de séries temporais, Brasil, 2007-2025”, e a "
     "expressão foi eliminada de todas as seções. O descritor “Estudos "
     "Ecológicos” foi substituído por “Estudos de Séries Temporais”, termo "
     "exato da DeCS. A "
     "ressalva sobre inferência ecológica permanece na discussão e na nota da "
     "Tabela 2, agora como cautela de interpretação da unidade de observação, "
     "e não como rótulo de delineamento."),

    ("Intervalos de confiança separados por ponto e vírgula.",
     "Todos os intervalos do resumo, do texto, das tabelas e das notas de "
     "ilustração passaram à forma “(IC95% 1,20; 1,26)”."),

    ("Retirar siglas não consagradas ou pouco usadas e explicar as consagradas "
     "na primeira menção de cada parte independente.",
     "Foram eliminadas SIH, AIH, SIM, OR, pp e ICr95%, substituídas pelos "
     "termos por extenso. As consagradas que permanecem — SINAN, CID-10, "
     "IC95%, STROBE e RECORD — são explicadas na primeira menção do resumo, na "
     "primeira menção do texto principal e no título ou nas notas de cada "
     "ilustração em que aparecem."),

    ("Declaração obrigatória de disponibilidade de dados, com link do "
     "repositório citado no texto e referência na lista.",
     "A seção “Disponibilidade de dados” nomeia as fontes públicas, a data das "
     "extrações e a conferência por resumo criptográfico, e informa o endereço "
     "do pacote de reprodutibilidade, citado no corpo e incluído como "
     "referência 15. O repositório é público e reúne os códigos de extração, "
     "decodificação e análise, o banco analítico derivado no nível região de "
     "saúde por ano, o dicionário de variáveis, a lista de verificação "
     "STROBE/RECORD, o manifesto de reprodutibilidade e o registro de "
     "rastreabilidade de cada valor publicado: " + REPO + " (etiqueta "
     "submissao-ress para a versão correspondente a esta submissão)."),

    ("Títulos de figuras e tabelas autossuficientes, com as siglas essenciais "
     "no próprio título.",
     "Os cinco títulos foram reescritos para se sustentarem sozinhos, com as "
     "siglas essenciais no título e, após ponto seguido, o local, os anos e o "
     "total de unidades incluídas na ilustração."),

    ("Conferir o espaçamento da Tabela 2.",
     "A tabela foi recomposta: espaçamento uniforme entre linhas, recuo "
     "interno homogêneo e linhas que não se partem entre páginas."),

    ("Sem subtítulos em Resultados e Discussão.",
     "Os subtítulos e os destaques em negrito no início de parágrafo foram "
     "removidos das duas seções, que passaram a texto corrido. Métodos "
     "conserva os tópicos previstos no guia de redação, na ordem preconizada."),

    ("As células das tabelas devem ter a mesma natureza declarada no cabeçalho.",
     "A Tabela 2 passou a quatro colunas de natureza única — Modelo, "
     "Especificação, Razão de chances (intervalo de 95%) e Células (n) —, e a "
     "antiga coluna de observações, que misturava naturezas, migrou para as "
     "notas. Na Tabela 1, os cabeçalhos passaram a declarar a estatística "
     "apresentada, como em “Incidência por 100.000 pessoas-ano (IC95%)”."),

    ("Referências em Vancouver, com periódico abreviado, página final "
     "abreviada ou e-page, link DOI, todos os endereços conferidos com data, e "
     "literatura cinzenta preferencialmente substituída por artigos.",
     "As 15 referências foram reformatadas em Vancouver, com periódico "
     "abreviado, página final abreviada ou e-page e link DOI. Todos os "
     "endereços foram resolvidos em 7 de setembro de 2026, data informada em "
     "cada referência; a citação de Clopper e Pearson passou do identificador "
     "do repositório para o do editor, ao qual aquele redireciona. Duas fontes "
     "de literatura cinzenta permanecem por não terem substituto em artigo "
     "científico: as projeções populacionais do Instituto Brasileiro de "
     "Geografia e Estatística, que são o denominador do estudo, e o Guia de "
     "Vigilância em Saúde, que é a definição de caso adotada. Ambas passaram a "
     "trazer endereço público conferido, e a segunda, o volume e as páginas em "
     "que a leptospirose é tratada."),
]


def main() -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    for lado in ("left", "right", "top", "bottom"):
        setattr(sec, "%s_margin" % lado, Cm(2.5))
    est = doc.styles["Normal"]
    est.font.name = FONTE
    est.font.size = CORPO

    p(doc, "**Carta de resposta — RESS-2026-1462**", espaco_depois=10)
    p(doc, "**Título:** Amplitude de detecção e letalidade notificada da "
           "leptospirose: análise de séries temporais, Brasil, 2007-2025")
    p(doc, "**Modalidade:** Artigo original", espaco_depois=12)

    p(doc,
      "Agradecemos a leitura atenta da secretaria. Todas as onze solicitações "
      "foram atendidas e estão relacionadas abaixo, na ordem em que foram "
      "numeradas, com o que se fez em cada uma. O manuscrito foi remontado por "
      "inteiro a partir do modelo da revista, de modo que as alterações "
      "alcançam também a formatação do documento.", espaco_depois=12)

    for i, (pedido, feito) in enumerate(ITENS, 1):
        p(doc, "**%d. %s**" % (i, pedido), espaco_antes=6, espaco_depois=3)
        p(doc, feito, espaco_depois=8, recuo=0.5)

    p(doc, "**Uma consulta sobre o sigilo da avaliação**", espaco_antes=10,
      espaco_depois=3)
    p(doc,
      "O item 6 pede o endereço do repositório citado no corpo do manuscrito e "
      "incluído na lista de referências, e foi assim que o atendemos. Como o "
      "endereço e a referência identificam a autoria, e a revista adota "
      "avaliação por pares cega, submetemos a questão à secretaria: se "
      "preferirem, substituímos ambos por uma versão anônima do repositório "
      "durante a avaliação e informamos o endereço definitivo na aceitação.",
      espaco_depois=8, recuo=0.5)

    doc.save(SAIDA)
    print("gravado:", SAIDA)


if __name__ == "__main__":
    main()
